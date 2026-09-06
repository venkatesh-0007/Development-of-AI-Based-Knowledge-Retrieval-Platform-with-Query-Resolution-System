"""Document extraction for PDF, DOCX, TXT and CSV files."""
from pathlib import Path
import io
import csv

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}

def extract_text(file_bytes: bytes, filename: str) -> str:
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    if ext == ".txt":
        return file_bytes.decode("utf-8", errors="replace")

    if ext == ".csv":
        try:
            import pandas as pd
            df = pd.read_csv(io.BytesIO(file_bytes))
            return df.to_csv(index=False)
        except ImportError:
            decoded = file_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(decoded))
            return "\n".join(", ".join(row) for row in reader)

    if ext == ".pdf":
        try:
            import fitz
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            return "\n".join(page.get_text("text") for page in doc)
        except ImportError:
            raise ImportError("PyMuPDF (fitz) is required to process PDF files. Please install requirements.txt.")

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))
            parts = [p.text for p in doc.paragraphs]
            for table in doc.tables:
                for row in table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
            return "\n".join(parts)
        except ImportError:
            raise ImportError("python-docx is required to process DOCX files. Please install requirements.txt.")
