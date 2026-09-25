"""Document extraction for PDF, DOCX, TXT and CSV files with section and domain metadata support."""
from pathlib import Path
import io
import csv
from typing import List, Dict, Any

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}

def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract raw text as a single string."""
    sections = extract_document_sections(file_bytes, filename)
    return "\n\n".join(s["text"] for s in sections if s.get("text"))

def extract_document_sections(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """
    Extract document text preserved as structured sections/pages.
    Returns: List[Dict] with 'text', 'page', 'section', and 'domain_hint'.
    """
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    lower_fn = filename.lower()
    if any(k in lower_fn for k in ["ml", "machine_learning", "deep_learning", "neural", "transformer", "algorithm"]):
        domain_hint = "machine_learning"
    elif any(k in lower_fn for k in ["network", "tcp", "udp", "protocol", "osi", "cloud", "distributed"]):
        domain_hint = "computer_networks"
    elif any(k in lower_fn for k in ["cyber", "security", "crypto", "zero_trust", "incident", "threat", "mitre", "cia"]):
        domain_hint = "cybersecurity"
    else:
        domain_hint = "general"

    if ext == ".txt":
        content = file_bytes.decode("utf-8", errors="replace")
        return [{"text": content, "page": None, "section": "Document Body", "domain_hint": domain_hint}]

    if ext == ".csv":
        try:
            import pandas as pd
            df = pd.read_csv(io.BytesIO(file_bytes))
            content = df.to_csv(index=False)
        except ImportError:
            decoded = file_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(decoded))
            content = "\n".join(", ".join(row) for row in reader)
        return [{"text": content, "page": None, "section": "Tabular Data", "domain_hint": domain_hint}]

    if ext == ".pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            sections = []
            for i, page in enumerate(doc):
                page_text = page.get_text("text").strip()
                if page_text:
                    sections.append({
                        "text": page_text,
                        "page": i + 1,
                        "section": f"Page {i + 1}",
                        "domain_hint": domain_hint
                    })
            if not sections:
                sections.append({"text": "", "page": 1, "section": "Page 1", "domain_hint": domain_hint})
            return sections
        except ImportError:
            raise ImportError("PyMuPDF (fitz) is required to process PDF files. Please install requirements.txt.")

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        parts.append(row_text)
            content = "\n".join(parts)
            return [{"text": content, "page": None, "section": "Document Body", "domain_hint": domain_hint}]
        except ImportError:
            raise ImportError("python-docx is required to process DOCX files. Please install requirements.txt.")

    return [{"text": "", "page": None, "section": "Document Body", "domain_hint": domain_hint}]
