"""Document extraction for PDF, DOCX, TXT and CSV files with section/page metadata support."""
from pathlib import Path
import io
import csv
from typing import List, Dict, Any

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".csv"}

def extract_text(file_bytes: bytes, filename: str) -> str:
    """Extract raw text as a single string (legacy/simple loader)."""
    sections = extract_document_sections(file_bytes, filename)
    return "\n\n".join(s["text"] for s in sections if s.get("text"))

def extract_document_sections(file_bytes: bytes, filename: str) -> List[Dict[str, Any]]:
    """
    Extract document text preserved as structured sections/pages.
    Returns a list of dicts: [{"text": str, "page": Optional[int], "section": str}]
    """
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"Unsupported file type: {ext}")

    if ext == ".txt":
        content = file_bytes.decode("utf-8", errors="replace")
        return [{"text": content, "page": None, "section": "Document Body"}]

    if ext == ".csv":
        try:
            import pandas as pd
            df = pd.read_csv(io.BytesIO(file_bytes))
            content = df.to_csv(index=False)
        except ImportError:
            decoded = file_bytes.decode("utf-8", errors="replace")
            reader = csv.reader(io.StringIO(decoded))
            content = "\n".join(", ".join(row) for row in reader)
        return [{"text": content, "page": None, "section": "Tabular Data"}]

    if ext == ".pdf":
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=file_bytes, filetype="pdf")
            sections = []
            for i, page in enumerate(doc):
                page_text = page.get_text("text").strip()
                if page_text:
                    sections.append({
                        "text": page_text,
                        "page": i + 1,
                        "section": f"Page {i + 1}"
                    })
            if not sections:
                sections.append({"text": "", "page": 1, "section": "Page 1"})
            return sections
        except ImportError:
            raise ImportError("PyMuPDF (fitz) is required to process PDF files. Please install requirements.txt.")

    if ext == ".docx":
        try:
            from docx import Document
            doc = Document(io.BytesIO(file_bytes))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        parts.append(row_text)
            content = "\n".join(parts)
            return [{"text": content, "page": None, "section": "Document Body"}]
        except ImportError:
            raise ImportError("python-docx is required to process DOCX files. Please install requirements.txt.")

    return [{"text": "", "page": None, "section": "Document Body"}]
